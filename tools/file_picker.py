#!/usr/bin/env python3
"""Local browser UI for selecting source files and running the auditor."""

from __future__ import annotations

import base64
import csv
import json
import os
import re
import shutil
import sys
import uuid
import webbrowser
from collections import OrderedDict
from html import escape
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import quote, unquote, urlparse

from openpyxl import load_workbook

WORKSPACE = Path(__file__).resolve().parents[1]
ENGINE_SCRIPTS = WORKSPACE / "multi-source-consistency-auditor" / "scripts"
if not ENGINE_SCRIPTS.is_dir():
    raise SystemExit("未找到核验引擎脚本目录")
sys.path.insert(0, str(ENGINE_SCRIPTS))

from multi_source_consistency_auditor.parsers import _document_chunks  # noqa: E402
from multi_source_consistency_auditor.profile import load_profile  # noqa: E402
from multi_source_consistency_auditor.reports import write_reports  # noqa: E402
from multi_source_consistency_auditor.rules import AuditEngine  # noqa: E402

HOST = "127.0.0.1"
PORT = 8780
RUNS = WORKSPACE / "user_data" / "file-picker-runs"
UPLOADS = WORKSPACE / "user_data" / "file-picker-uploads"
ALLOWED_REPORTS = {
    "differences-summary.pdf",
    "differences-summary.docx",
    "differences-summary.xlsx",
    "differences-summary.html",
    "differences-summary.md",
    "differences-summary.json",
    "differences-summary.csv",
    "profile.json",
}

KNOWN_FIELDS = [
    {
        "label": "名称",
        "aliases": ["名称", "项目名称", "采购项目", "合同名称", "标题", "事项"],
        "type": "text",
    },
    {
        "label": "供应商",
        "aliases": ["供应商", "供应商名称", "乙方", "服务商", "供方"],
        "type": "text",
    },
    {
        "label": "金额",
        "aliases": ["金额", "合同金额", "采购金额", "报价金额", "总金额", "价税合计"],
        "type": "amount",
    },
    {
        "label": "数量",
        "aliases": ["数量", "采购数量", "件数", "人数"],
        "type": "number",
    },
]

PAGE = """<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>选择文件并核验</title>
<style>
:root{--ink:#17212b;--muted:#607080;--line:#d6e0e7;--teal:#156f73;--dark:#12384b;--bg:#f2f5f7}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--ink);font-family:"Microsoft YaHei",sans-serif;line-height:1.6}
main{max-width:900px;margin:5vh auto;padding:0 18px 40px}.card{background:white;border:1px solid var(--line);border-radius:16px;padding:28px;box-shadow:0 14px 36px #12384b14}
h1{color:var(--dark);margin:0 0 8px}.lead{color:var(--muted);margin-top:0}.zone{border:2px dashed #9fb8bd;border-radius:12px;padding:28px;text-align:center;background:#f8fbfb;margin:22px 0}
input[type=file]{display:none}.pick{display:inline-block;background:var(--teal);color:white;padding:11px 22px;border-radius:8px;font-weight:700;cursor:pointer}
.grid{display:grid;grid-template-columns:1fr 1fr;gap:14px}.field{margin-top:10px}.field label{font-weight:700;display:block;margin-bottom:6px}
input[type=text]{width:100%;padding:10px;border:1px solid #b8c7d1;border-radius:7px;font-size:15px}
.options{display:flex;flex-wrap:wrap;gap:12px 22px;margin:12px 0}.options label{cursor:pointer}
button{background:var(--teal);color:white;border:0;border-radius:8px;padding:12px 24px;font-size:16px;font-weight:700;cursor:pointer;margin-top:18px}
button:disabled{opacity:.55;cursor:wait}.note{background:#f4f8f9;border-left:4px solid var(--teal);padding:12px 15px;color:#526575;font-size:13px}
.files{margin-top:14px;color:var(--muted);font-size:13px}.file-row{display:flex;align-items:center;justify-content:space-between;gap:12px;background:#fff;border:1px solid var(--line);border-radius:7px;padding:7px 10px;margin-top:7px;color:var(--ink)}.file-row span{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}.remove{border:1px solid #d7a9a4;background:#fff5f3;color:#9b2920;padding:4px 9px;font-size:12px;margin:0}.clear{border:0;background:transparent;color:var(--teal);padding:5px;margin-top:8px;font-size:13px}.result{margin-top:22px;padding:18px;border-radius:10px;background:#f5f9f9;border:1px solid var(--line)}
.links a{display:inline-block;margin:7px 12px 0 0;color:var(--teal);font-weight:700}.error{background:#fff0ee;border-color:#f1b9b3;color:#9b2920}
@media(max-width:650px){.grid{grid-template-columns:1fr}main{margin-top:20px}}
</style>
</head>
<body><main><div class="card">
<h1>选择文件并核验</h1>
<p class="lead">一次选择多份表格和文档。程序只读取本地副本，不修改、重命名或转换原文件。</p>
<div class="zone">
<input id="files" type="file" multiple accept=".csv,.tsv,.xlsx,.xlsm,.json,.docx,.pdf,.txt,.md,.markdown,.html,.htm">
<label class="pick" for="files">选择文件</label>
<div class="files">可以多次点击“选择文件”，每次选择都会追加到下方列表。</div>
<div id="file-list" class="files">尚未选择文件</div>
<button id="clear-files" class="clear" type="button" hidden>清空全部已选文件</button>
</div>
<div class="grid">
<div class="field"><label for="id-field">统一编号字段</label><input id="id-field" type="text" value="编号" placeholder="例如：合同编号、项目编号、学号"></div>
<div class="field"><label>核验选项</label><label><input id="coverage" type="checkbox" checked> 检查每个文件是否包含同一编号</label></div>
</div>
<div class="field"><label>输出格式</label><div class="options">
<label><input type="checkbox" name="format" value="pdf" checked> PDF</label>
<label><input type="checkbox" name="format" value="docx" checked> Word</label>
<label><input type="checkbox" name="format" value="xlsx" checked> Excel</label>
<label><input type="checkbox" name="format" value="html"> HTML</label>
<label><input type="checkbox" name="format" value="json"> JSON</label>
</div></div>
<button id="run" type="button">开始核验</button>
<div class="note">表格优先按编号列关联；文档字段建议采用“字段名：值”的形式，例如“合同编号：P-2026-001”。如果编号不在正文，需要把编号写进文件名并由后续配置正则提取。</div>
<div id="result" class="result" hidden></div>
</div></main>
<script>
const input=document.getElementById("files"),list=document.getElementById("file-list"),button=document.getElementById("run"),result=document.getElementById("result"),clear=document.getElementById("clear-files");
let selectedFiles=[];
function fileKey(file){return `${file.name}:${file.size}:${file.lastModified}`;}
function renderFiles(){
  list.innerHTML="";
  if(!selectedFiles.length){list.textContent="尚未选择文件";clear.hidden=true;return;}
  clear.hidden=false;
  selectedFiles.forEach((file,index)=>{
    const row=document.createElement("div");row.className="file-row";
    const name=document.createElement("span");name.textContent=`${index+1}. ${file.name}`;
    const remove=document.createElement("button");remove.type="button";remove.className="remove";remove.dataset.index=String(index);remove.textContent="移除";
    row.append(name,remove);list.append(row);
  });
}
input.addEventListener("change",()=>{
  const existing=new Set(selectedFiles.map(fileKey));
  for(const file of [...input.files]){const key=fileKey(file);if(!existing.has(key)){selectedFiles.push(file);existing.add(key);}}
  input.value="";renderFiles();
});
list.addEventListener("click",(event)=>{
  const remove=event.target.closest("button[data-index]");if(!remove)return;
  selectedFiles.splice(Number(remove.dataset.index),1);renderFiles();
});
clear.addEventListener("click",()=>{selectedFiles=[];renderFiles();});
function readFile(file){return new Promise((resolve,reject)=>{const reader=new FileReader();reader.onload=()=>resolve({name:file.name,data:reader.result});reader.onerror=reject;reader.readAsDataURL(file);});}
button.addEventListener("click",async()=>{
  const files=[...selectedFiles];if(!files.length){alert("请先选择文件");return;}
  const formats=[...document.querySelectorAll('input[name="format"]:checked')].map(x=>x.value);
  if(!formats.length){alert("请至少选择一种输出格式");return;}
  button.disabled=true;button.textContent="正在核验…";result.hidden=true;
  try{
    const payload={files:await Promise.all(files.map(readFile)),id_field:document.getElementById("id-field").value.trim()||"编号",coverage:document.getElementById("coverage").checked,formats};
    const response=await fetch("/run",{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(payload)});
    const data=await response.json();
    if(!response.ok) throw new Error(data.error||"核验失败");
    const links=Object.entries(data.links||{}).map(([name,url])=>`<a href="${url}" target="_blank">${name}</a>`).join("");
    result.className="result";result.innerHTML=`<strong>核验完成：${data.status}</strong><br>实体 ${data.entities}，通过 ${data.passed}，错误 ${data.errors}，警告 ${data.warnings}<div class="links">${links}</div><p class="files">报告目录：${data.run_folder}<br>本次文件副本：${data.upload_folder}</p>`;result.hidden=false;
  }catch(error){result.className="result error";result.innerHTML=`<strong>核验失败</strong><br>${String(error)}`;result.hidden=false;}
  finally{button.disabled=false;button.textContent="开始核验";}
});
</script>
</body></html>"""


def clean_filename(name: str, index: int) -> str:
    base = Path(name).name
    base = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", base)
    if not base:
        base = f"file-{index}"
    return f"{index:03d}_{base}"


def normalize_label(value: str) -> str:
    text = re.sub(r"[\s:：()（）_\-]+", "", value or "")
    return text.casefold()


def detect_type(label: str, value: str = "") -> str:
    if any(token in label for token in ("金额", "价格", "费用", "合计")):
        return "amount"
    if any(token in label for token in ("日期", "时间")):
        return "date"
    if any(token in label for token in ("数量", "人数", "件数")):
        return "number"
    return "text"


def table_labels(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix in {".csv", ".tsv"}:
        for encoding in ("utf-8-sig", "gb18030"):
            try:
                with path.open("r", encoding=encoding, newline="") as handle:
                    sample = handle.read(4096)
                    handle.seek(0)
                    delimiter = "\t" if suffix == ".tsv" else ","
                    if suffix == ".csv":
                        try:
                            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
                        except csv.Error:
                            pass
                    return [item.strip() for item in next(csv.reader(handle, delimiter=delimiter), []) if item.strip()]
            except UnicodeDecodeError:
                continue
        return []
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if len(values) >= 2:
                        return values
        finally:
            workbook.close()
        return []
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        rows = data if isinstance(data, list) else data.get("records", []) if isinstance(data, dict) else []
        return list(rows[0].keys()) if rows and isinstance(rows[0], dict) else []
    return []


def document_labels(path: Path) -> list[str]:
    labels: OrderedDict[str, str] = OrderedDict()
    for _location, text in _document_chunks(path):
        for segment in re.split(r"[\n\r；;|]+", text):
            match = re.match(r"^\s*([^:：]{1,40})\s*[:：]\s*(.+?)\s*$", segment)
            if match:
                label = match.group(1).strip()
                if label and label not in labels:
                    labels[label] = match.group(2).strip()
    return list(labels.keys())


def build_profile(
    session: Path,
    upload_paths: list[Path],
    id_field: str,
    coverage: bool,
) -> tuple[Path, int]:
    labels: OrderedDict[str, str] = OrderedDict()
    for path in upload_paths:
        source_labels = table_labels(path) if path.suffix.lower() in {".csv", ".tsv", ".xlsx", ".xlsm", ".json"} else document_labels(path)
        for label in source_labels:
            labels.setdefault(normalize_label(label), label)

    id_key = normalize_label(id_field)
    field_specs: OrderedDict[str, dict[str, object]] = OrderedDict()
    for label_key, label in labels.items():
        if label_key == id_key or any(alias in label_key or label_key in alias for alias in (id_key,)):
            continue
        matched = None
        for known in KNOWN_FIELDS:
            if any(normalize_label(alias) in label_key or label_key in normalize_label(alias) for alias in known["aliases"]):
                matched = known
                break
        key = ""
        if matched:
            for existing_key, existing in field_specs.items():
                if existing.get("_group") == matched["label"]:
                    key = existing_key
                    break
            if not key:
                key = f"field_{len(field_specs) + 1}"
                field_specs[key] = {
                    "label": matched["label"],
                    "aliases": list(matched["aliases"]),
                    "type": matched["type"],
                    "_group": matched["label"],
                }
            aliases = field_specs[key]["aliases"]
            if isinstance(aliases, list) and label not in aliases:
                aliases.append(label)
        else:
            key = f"field_{len(field_specs) + 1}"
            field_specs[key] = {
                "label": label,
                "aliases": [label],
                "type": detect_type(label),
                "_group": label,
            }

    if not field_specs:
        raise ValueError("没有识别到可核验字段。请确认表格有列名，文档采用“字段名：值”格式。")
    for spec in field_specs.values():
        spec.pop("_group", None)

    identifier_aliases = []
    for value in [id_field, "编号", "项目编号", "合同编号", "订单号", "ID"]:
        if value and value not in identifier_aliases:
            identifier_aliases.append(value)

    sources = []
    for index, path in enumerate(upload_paths, start=1):
        source_type = "table" if path.suffix.lower() in {".csv", ".tsv", ".xlsx", ".xlsm", ".json"} else "document"
        sources.append(
            {
                "name": f"{index:02d} {path.stem}",
                "type": source_type,
                "paths": [str(path)],
                "expected": coverage,
            }
        )

    checks = [
        {
            "type": "consistent",
            "field": key,
            "severity": "warning",
        }
        for key in field_specs
    ]
    profile = {
        "name": f"{session.name} 文件核验",
        "entity": "业务实体",
        "identifier": {"field": "entity_id", "aliases": identifier_aliases},
        "sources": sources,
        "fields": field_specs,
        "tolerances": {
            key: "0.01" for key, spec in field_specs.items() if spec["type"] == "amount"
        },
        "auto_consistent_fields": False,
        "coverage_severity": "warning",
        "checks": checks,
    }
    profile_path = session / "profile.json"
    profile_path.write_text(json.dumps(profile, ensure_ascii=False, indent=2), encoding="utf-8")
    return profile_path, len(field_specs)


class Handler(BaseHTTPRequestHandler):
    def log_message(self, format: str, *args: object) -> None:
        return

    def send_content(self, status: int, content_type: str, body: bytes) -> None:
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def do_GET(self) -> None:
        parsed = urlparse(self.path)
        if parsed.path == "/":
            self.send_content(200, "text/html; charset=utf-8", PAGE.encode("utf-8"))
            return
        match = re.fullmatch(r"/runs/([a-f0-9]{32})/([^/]+)", parsed.path)
        if match and match.group(2) in ALLOWED_REPORTS:
            path = RUNS / match.group(1) / match.group(2)
            if path.is_file():
                content_types = {
                    ".pdf": "application/pdf",
                    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ".html": "text/html; charset=utf-8",
                    ".json": "application/json; charset=utf-8",
                    ".md": "text/markdown; charset=utf-8",
                    ".csv": "text/csv; charset=utf-8",
                }
                self.send_content(200, content_types[path.suffix.lower()], path.read_bytes())
                return
        self.send_content(404, "text/plain; charset=utf-8", "Not Found".encode("utf-8"))

    def do_POST(self) -> None:
        if urlparse(self.path).path != "/run":
            self.send_content(404, "text/plain; charset=utf-8", b"Not Found")
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0 or length > 150 * 1024 * 1024:
                raise ValueError("提交内容为空或超过 150 MB 限制")
            payload = json.loads(self.rfile.read(length).decode("utf-8"))
            raw_files = payload.get("files", [])
            if not raw_files or len(raw_files) > 50:
                raise ValueError("请选择 1—50 个文件")
            formats = [item for item in payload.get("formats", []) if item in {"pdf", "docx", "xlsx", "html", "json", "md", "csv"}]
            if not formats:
                raise ValueError("请至少选择一种输出格式")
            run_id = uuid.uuid4().hex
            upload_dir = UPLOADS / run_id
            run_dir = RUNS / run_id
            upload_dir.mkdir(parents=True)
            run_dir.mkdir(parents=True)
            uploaded: list[Path] = []
            allowed_suffixes = {".csv", ".tsv", ".xlsx", ".xlsm", ".json", ".docx", ".pdf", ".txt", ".md", ".markdown", ".html", ".htm"}
            total_bytes = 0
            for index, item in enumerate(raw_files, start=1):
                filename = clean_filename(str(item.get("name", "")), index)
                suffix = Path(filename).suffix.lower()
                if suffix not in allowed_suffixes:
                    raise ValueError(f"不支持的文件类型：{suffix or '无扩展名'}")
                data_url = str(item.get("data", ""))
                if "," not in data_url:
                    raise ValueError(f"文件内容无效：{filename}")
                content = base64.b64decode(data_url.split(",", 1)[1], validate=True)
                total_bytes += len(content)
                if total_bytes > 120 * 1024 * 1024:
                    raise ValueError("所选文件总大小超过 120 MB")
                path = upload_dir / filename
                path.write_bytes(content)
                uploaded.append(path)
            profile_path, field_count = build_profile(
                run_dir,
                uploaded,
                str(payload.get("id_field") or "编号"),
                bool(payload.get("coverage", True)),
            )
            result = AuditEngine(load_profile(profile_path)).run()
            report_paths = write_reports(result, run_dir, formats)
            summary = result.to_dict()["summary"]
            links = {
                name: f"/runs/{run_id}/{quote(path.name)}"
                for name, path in report_paths.items()
            }
            response = {
                "status": result.status,
                "entities": summary["entities"],
                "passed": summary["passed_entities"],
                "errors": summary["errors"],
                "warnings": summary["warnings"],
                "fields": field_count,
                "links": links,
                "run_folder": str(run_dir),
                "upload_folder": str(upload_dir),
            }
            self.send_content(200, "application/json; charset=utf-8", json.dumps(response, ensure_ascii=False).encode("utf-8"))
        except Exception as exc:
            self.send_content(400, "application/json; charset=utf-8", json.dumps({"error": str(exc)}, ensure_ascii=False).encode("utf-8"))


def main() -> None:
    UPLOADS.mkdir(parents=True, exist_ok=True)
    RUNS.mkdir(parents=True, exist_ok=True)
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    url = f"http://{HOST}:{PORT}/"
    print(f"文件选择核验工具已启动：{url}")
    print("使用期间请不要关闭本窗口；按 Ctrl+C 可以停止。")
    if os.environ.get("MSCA_NO_BROWSER") != "1":
        try:
            if not webbrowser.open(url):
                print(f"浏览器未能自动打开，请手动访问：{url}")
        except Exception as exc:
            print(f"浏览器未能自动打开，请手动访问：{url}（{exc}）")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()


if __name__ == "__main__":
    main()
