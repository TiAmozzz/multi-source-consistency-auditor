#!/usr/bin/env python3
"""Rebuild the fictional procurement sample documents."""

from __future__ import annotations

import csv
from pathlib import Path

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfgen import canvas

BASE = Path(__file__).resolve().parent / "procurement" / "input"
CONTRACT_DIR = BASE / "合同文件"


def build_spreadsheet() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "采购记录"
    rows = [
        ["项目编号", "项目名称", "供应商", "采购金额", "申请日期"],
        ["P-2026-001", "办公电脑采购", "供应商甲", 12000, "2026-09-01"],
        ["P-2026-002", "激光打印机采购", "供应商乙", 3500, "2026-09-15"],
        ["P-2026-003", "网络设备采购", "供应商丙", 8000, "2026-09-20"],
    ]
    for row in rows:
        sheet.append(row)
    workbook.save(BASE / "采购台账.xlsx")


def build_csv() -> None:
    rows = [
        ["采购编号", "采购项目", "供应商名称", "合同金额", "申请日期"],
        ["P-2026-001", "办公电脑采购", "供应商甲", "12,000.00", "2026-09-01"],
        ["P-2026-002", "激光打印机采购", "供应商乙", "3,200.00", "2026-09-15"],
        ["P-2026-004", "会议扩声设备采购", "供应商丁", "4,500.00", "2026-09-21"],
    ]
    with (BASE / "采购申请.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        csv.writer(handle).writerows(rows)


def build_docx(
    filename: str,
    project_id: str,
    title: str,
    supplier: str,
    amount: str,
    signed_date: str,
) -> None:
    document = Document()
    document.add_heading("采购合同", level=1)
    document.add_paragraph("本文件仅用于开源项目功能演示，不具备真实合同效力。")
    document.add_paragraph(f"项目编号：{project_id}")
    document.add_paragraph(f"项目名称：{title}")
    document.add_paragraph(f"供应商：{supplier}")
    document.add_paragraph(f"合同金额：人民币 {amount} 元")
    document.add_paragraph(f"签订日期：{signed_date}")
    document.save(CONTRACT_DIR / filename)


def build_pdf() -> None:
    path = CONTRACT_DIR / "P-2026-005_备用电源采购合同.pdf"
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf = canvas.Canvas(str(path), pagesize=A4)
    pdf.setFont("STSong-Light", 12)
    lines = [
        "采购合同（演示文件）",
        "",
        "项目编号：P-2026-005",
        "项目名称：备用电源采购",
        "供应商：供应商戊",
        "合同金额：人民币 6,000.00 元",
        "签订日期：2026年9月25日",
        "",
        "本文件仅用于开源项目功能演示，不具备真实合同效力。",
    ]
    height = A4[1] - 72
    for line in lines:
        pdf.drawString(72, height, line)
        height -= 24
    pdf.save()


def main() -> None:
    BASE.mkdir(parents=True, exist_ok=True)
    CONTRACT_DIR.mkdir(parents=True, exist_ok=True)
    build_spreadsheet()
    build_csv()
    build_docx(
        "P-2026-001_办公电脑采购合同.docx",
        "P-2026-001",
        "办公电脑采购",
        "供应商甲",
        "12,000.00",
        "2026年9月8日",
    )
    build_docx(
        "P-2026-002_激光打印机采购合同.docx",
        "P-2026-002",
        "激光打印机采购",
        "供应商乙",
        "3,500.00",
        "2026年9月19日",
    )
    build_pdf()
    print(f"已生成虚构样例：{BASE}")


if __name__ == "__main__":
    main()
