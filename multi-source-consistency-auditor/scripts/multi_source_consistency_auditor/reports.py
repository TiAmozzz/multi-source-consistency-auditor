"""Generate human-readable and machine-readable audit reports."""

from __future__ import annotations

import csv
import html
import json
import os
from pathlib import Path
from typing import Any, Iterable

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from xml.sax.saxutils import escape

from . import __version__
from .model import AuditResult, Issue


def _safe_text(value: Any) -> str:
    text = "" if value is None else str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _safe_csv(path: Path, headers: list[str], rows: Iterable[Iterable[Any]]) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            writer.writerow([_safe_text(value) for value in row])


def _evidence_text(issue: Issue) -> str:
    return "\n".join(
        f"{item.file} / {item.location}: {item.excerpt}" for item in issue.evidence
    )


def _markdown_escape(value: Any) -> str:
    return str(value or "").replace("|", "\\|").replace("\n", "<br>")


def _entity_values(entity: dict[str, Any], labels: dict[str, str]) -> str:
    parts: list[str] = []
    for field_name, values in entity.get("fields", {}).items():
        label = labels.get(field_name, field_name)
        rendered = "；".join(f"{item['value']} [{item['source']}]" for item in values)
        parts.append(f"{label}：{rendered}")
    return "；".join(parts)


def _write_markdown(result: AuditResult, path: Path) -> None:
    summary = result.to_dict()["summary"]
    lines = [
        f"# {result.profile_name}报告",
        "",
        f"- 核验对象：{result.entity_name}",
        f"- 总体状态：**{result.status}**",
        f"- 来源数量：{summary['sources']}",
        f"- 文件数量：{summary['files']}",
        f"- 实体数量：{summary['entities']}",
        f"- 错误：{summary['errors']}",
        f"- 警告：{summary['warnings']}",
        f"- 输入文件完整性：{'未改动' if not result.modified_input_files else '检测到改动'}",
        f"- 生成时间：{result.finished_at}",
        "",
        "## 实体概览",
        "",
        "| 编号 | 状态 | 来源 | 问题数 |",
        "|---|---|---|---:|",
    ]
    for entity in result.entity_summaries:
        lines.append(
            f"| {_markdown_escape(entity['entity_id'])} | {entity['status']} | "
            f"{_markdown_escape('、'.join(entity['sources']))} | {entity['issue_count']} |"
        )

    lines.extend(
        [
            "",
            "## 问题明细",
            "",
            "| 严重程度 | 编号 | 实体 | 字段 | 说明 |",
            "|---|---|---|---|---|",
        ]
    )
    if not result.issue_list:
        lines.append("| - | - | - | - | 未发现问题 |")
    for issue in result.issue_list:
        lines.append(
            f"| {issue.severity} | {_markdown_escape(issue.code)} | "
            f"{_markdown_escape(issue.entity_id)} | "
            f"{_markdown_escape(issue.field_name or '')} | "
            f"{_markdown_escape(issue.message)} |"
        )

    lines.extend(["", "## 证据索引", ""])
    for issue in result.issue_list:
        if not issue.evidence:
            continue
        lines.append(
            f"### {_markdown_escape(issue.entity_id)} / {_markdown_escape(issue.code)}"
        )
        lines.append("")
        lines.append(f"结论：{issue.message}")
        for item in issue.evidence:
            lines.append(
                f"- `{_markdown_escape(item.file)}` / {_markdown_escape(item.location)}："
                f"{_markdown_escape(item.excerpt)}"
            )
        lines.append("")

    lines.extend(
        [
            "",
            "## 使用限制",
            "",
            "本报告只反映配置规则的核验结果，不判断文件内容在业务或法律意义上的真实性。"
            "带警告的项目应由人工结合原始资料复核。",
            "",
        ]
    )
    path.write_text("\n".join(lines), encoding="utf-8")


def _style_sheet(sheet: Any, widths: dict[int, int] | None = None) -> None:
    header_fill = PatternFill("solid", fgColor="16324F")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    if widths:
        for index, width in widths.items():
            sheet.column_dimensions[get_column_letter(index)].width = width


def _write_workbook(result: AuditResult, path: Path) -> None:
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "汇总"
    summary = result.to_dict()["summary"]
    summary_rows = [
        ("核验配置", result.profile_name),
        ("核验对象", result.entity_name),
        ("总体状态", result.status),
        ("来源数量", summary["sources"]),
        ("文件数量", summary["files"]),
        ("记录数量", summary["records"]),
        ("实体数量", summary["entities"]),
        ("通过实体", summary["passed_entities"]),
        ("需复核实体", summary["review_entities"]),
        ("错误", summary["errors"]),
        ("警告", summary["warnings"]),
        (
            "输入文件完整性",
            "未改动" if not result.modified_input_files else "检测到改动",
        ),
        ("开始时间", result.started_at),
        ("结束时间", result.finished_at),
    ]
    summary_sheet.append(["项目", "内容"])
    for row in summary_rows:
        summary_sheet.append(list(row))
    _style_sheet(summary_sheet, {1: 18, 2: 48})

    issues_sheet = workbook.create_sheet("问题明细")
    issues_sheet.append(
        [
            "严重程度",
            "问题代码",
            "实体编号",
            "字段",
            "问题说明",
            "期望",
            "实际",
            "证据",
        ]
    )
    for issue in result.issue_list:
        issues_sheet.append(
            [
                issue.severity,
                issue.code,
                _safe_text(issue.entity_id),
                issue.field_name or "",
                issue.message,
                _safe_text(issue.expected),
                _safe_text(issue.actual),
                _evidence_text(issue),
            ]
        )
    _style_sheet(
        issues_sheet,
        {1: 12, 2: 25, 3: 24, 4: 18, 5: 42, 6: 34, 7: 42, 8: 64},
    )

    entity_sheet = workbook.create_sheet("实体对照")
    # Use the stable field order carried by each entity summary.
    labels: dict[str, str] = {}
    for entity in result.entity_summaries:
        labels.update(entity.get("labels", {}))
    field_headers = list(labels.values())
    entity_sheet.append(["实体编号", "状态", "来源", "问题数", *field_headers])
    label_items = list(labels.items())
    for entity in result.entity_summaries:
        values = []
        for field_name, _label in label_items:
            field_values = entity.get("fields", {}).get(field_name, [])
            values.append(
                _safe_text(
                    "；".join(
                        f"{item['value']} [{item['source']}]" for item in field_values
                    )
                )
            )
        entity_sheet.append(
            [
                _safe_text(entity["entity_id"]),
                entity["status"],
                "、".join(entity["sources"]),
                entity["issue_count"],
                *values,
            ]
        )
    _style_sheet(
        entity_sheet,
        {1: 24, 2: 12, 3: 34, 4: 10, **{i: 36 for i in range(5, 5 + len(field_headers))}},
    )
    if entity_sheet.max_column > 5:
        entity_sheet.auto_filter.ref = entity_sheet.dimensions

    evidence_sheet = workbook.create_sheet("证据索引")
    evidence_sheet.append(
        ["实体编号", "问题代码", "来源", "文件", "位置", "原文片段"]
    )
    for issue in result.issue_list:
        for item in issue.evidence:
            evidence_sheet.append(
                [
                    _safe_text(issue.entity_id),
                    issue.code,
                    item.source,
                    item.file,
                    item.location,
                    item.excerpt,
                ]
            )
    _style_sheet(evidence_sheet, {1: 24, 2: 24, 3: 18, 4: 42, 5: 24, 6: 80})

    source_sheet = workbook.create_sheet("来源清单")
    source_sheet.append(["来源", "文件", "类型", "状态", "记录数"])
    for item in result.source_inventory:
        source_sheet.append(
            [item.source, item.file, item.source_type, item.status, item.records]
        )
    _style_sheet(source_sheet, {1: 22, 2: 58, 3: 14, 4: 14, 5: 12})

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.save(path)


def _status_badge(status: str) -> str:
    class_name = "pass" if status == "通过" else "review" if status == "需复核" else "fail"
    return f'<span class="badge {class_name}">{html.escape(status)}</span>'


def _write_html(result: AuditResult, path: Path) -> None:
    payload = result.to_dict()
    summary = payload["summary"]
    labels: dict[str, str] = {}
    for entity in result.entity_summaries:
        labels.update(entity.get("labels", {}))

    entities_html: list[str] = []
    for entity in result.entity_summaries:
        values = []
        for field_name, values_list in entity.get("fields", {}).items():
            rendered = "<br>".join(
                f"{html.escape(item['value'])} "
                f'<span class="source">{html.escape(item["source"])}</span>'
                for item in values_list
            )
            values.append(
                f"<strong>{html.escape(labels.get(field_name, field_name))}</strong><br>{rendered}"
            )
        entities_html.append(
            "<tr>"
            f"<td><code>{html.escape(entity['entity_id'])}</code></td>"
            f"<td>{_status_badge(entity['status'])}</td>"
            f"<td>{html.escape('、'.join(entity['sources']))}</td>"
            f"<td>{'<br>'.join(values)}</td>"
            f"<td>{entity['issue_count']}</td>"
            "</tr>"
        )
    if not entities_html:
        entities_html.append('<tr><td colspan="5">未解析到实体</td></tr>')

    issues_html: list[str] = []
    for issue in result.issue_list:
        evidence = "".join(
            "<li>"
            f"<code>{html.escape(item.file)}</code> / {html.escape(item.location)}"
            f"<blockquote>{html.escape(item.excerpt)}</blockquote>"
            "</li>"
            for item in issue.evidence
        )
        evidence_block = (
            f"<details><summary>查看 {len(issue.evidence)} 条证据</summary><ul>{evidence}</ul></details>"
            if issue.evidence
            else ""
        )
        issues_html.append(
            f'<tr class="severity-{html.escape(issue.severity)}">'
            f"<td>{html.escape(issue.severity)}</td>"
            f"<td><code>{html.escape(issue.code)}</code></td>"
            f"<td><code>{html.escape(issue.entity_id)}</code></td>"
            f"<td>{html.escape(labels.get(issue.field_name or '', issue.field_name or ''))}</td>"
            f"<td><strong>{html.escape(issue.message)}</strong>"
            f"<br><small>期望：{html.escape(issue.expected or '未指定')}</small>"
            f"<br><small>实际：{html.escape(issue.actual or '未指定')}</small>"
            f"{evidence_block}</td></tr>"
        )
    if not issues_html:
        issues_html.append('<tr><td colspan="5">未发现问题</td></tr>')

    source_rows = "".join(
        "<tr>"
        f"<td>{html.escape(item.source)}</td>"
        f"<td><code>{html.escape(item.file)}</code></td>"
        f"<td>{html.escape(item.source_type)}</td>"
        f"<td>{html.escape(item.status)}</td>"
        f"<td>{item.records}</td>"
        "</tr>"
        for item in result.source_inventory
    )

    document = f"""<!doctype html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html.escape(result.profile_name)}核验报告</title>
<style>
:root {{ --ink:#17212b; --muted:#607080; --line:#dbe3ea; --paper:#fff; --bg:#f2f5f8; --accent:#156f73; }}
* {{ box-sizing:border-box; }}
body {{ margin:0; font-family:"Microsoft YaHei","Noto Sans SC",sans-serif; color:var(--ink); background:var(--bg); line-height:1.55; }}
main {{ max-width:1320px; margin:28px auto; padding:0 24px 48px; }}
.hero {{ background:linear-gradient(135deg,#12384b,#156f73); color:white; border-radius:18px; padding:30px 34px; box-shadow:0 12px 32px #12384b22; }}
.hero h1 {{ margin:0 0 8px; font-size:30px; }}
.hero p {{ margin:0; color:#dceff0; }}
.stats {{ display:grid; grid-template-columns:repeat(5,minmax(120px,1fr)); gap:12px; margin:20px 0; }}
.stat {{ background:var(--paper); border:1px solid var(--line); border-radius:12px; padding:16px; }}
.stat b {{ display:block; font-size:27px; color:var(--accent); }}
.stat span {{ color:var(--muted); font-size:13px; }}
section {{ background:var(--paper); border:1px solid var(--line); border-radius:14px; margin-top:18px; overflow:hidden; }}
section h2 {{ margin:0; padding:17px 20px; font-size:19px; border-bottom:1px solid var(--line); }}
table {{ width:100%; border-collapse:collapse; font-size:14px; }}
th {{ background:#f6f8fa; text-align:left; color:#40505e; font-weight:600; }}
th,td {{ padding:11px 13px; border-bottom:1px solid var(--line); vertical-align:top; }}
tr:last-child td {{ border-bottom:0; }}
code {{ background:#eef2f5; border-radius:4px; padding:1px 5px; }}
.badge {{ display:inline-block; padding:3px 9px; border-radius:999px; font-size:12px; font-weight:700; white-space:nowrap; }}
.pass {{ color:#0e7043; background:#dcf6e7; }}
.review {{ color:#8b5c00; background:#fff0c8; }}
.fail {{ color:#a52b2b; background:#ffe1df; }}
.source {{ color:var(--muted); font-size:11px; border:1px solid var(--line); border-radius:999px; padding:0 5px; }}
.severity-error td:first-child {{ color:#a52b2b; font-weight:700; }}
.severity-warning td:first-child {{ color:#8b5c00; font-weight:700; }}
details {{ margin-top:7px; }}
summary {{ cursor:pointer; color:var(--accent); }}
blockquote {{ margin:5px 0; padding:6px 10px; border-left:3px solid #9fb9bb; color:#42505d; background:#f7f9fa; }}
.note {{ padding:14px 20px; color:var(--muted); font-size:13px; }}
@media(max-width:900px) {{ .stats {{ grid-template-columns:repeat(2,1fr); }} table {{ font-size:12px; }} main {{ padding:0 10px 30px; }} }}
</style>
</head>
<body>
<main>
  <header class="hero">
    <h1>{html.escape(result.profile_name)}报告</h1>
    <p>核验对象：{html.escape(result.entity_name)} · 生成时间：{html.escape(result.finished_at)} · 状态：{_status_badge(result.status)}</p>
  </header>
  <div class="stats">
    <div class="stat"><b>{summary['files']}</b><span>解析文件</span></div>
    <div class="stat"><b>{summary['entities']}</b><span>关联实体</span></div>
    <div class="stat"><b>{summary['passed_entities']}</b><span>通过实体</span></div>
    <div class="stat"><b>{summary['errors']}</b><span>错误</span></div>
    <div class="stat"><b>{summary['warnings']}</b><span>警告</span></div>
  </div>
  <section>
    <h2>实体对照</h2>
    <table><thead><tr><th>实体编号</th><th>状态</th><th>已覆盖来源</th><th>字段对照</th><th>问题数</th></tr></thead>
    <tbody>{''.join(entities_html)}</tbody></table>
  </section>
  <section>
    <h2>问题明细</h2>
    <table><thead><tr><th>级别</th><th>代码</th><th>实体</th><th>字段</th><th>说明与证据</th></tr></thead>
    <tbody>{''.join(issues_html)}</tbody></table>
  </section>
  <section>
    <h2>来源清单</h2>
    <table><thead><tr><th>来源</th><th>文件</th><th>类型</th><th>解析状态</th><th>记录数</th></tr></thead>
    <tbody>{source_rows}</tbody></table>
  </section>
  <p class="note">本报告仅反映配置规则的计算结果。字段缺失、来源缺失和数值差异不等同于业务事实或法律结论，异常项应由责任人结合原始资料复核。</p>
  <p class="note">原始资料以只读方式解析；程序不转换、重命名或回写源文件，不联网上传资料。输入文件完整性：{'未改动' if not result.modified_input_files else '检测到改动'}。</p>
</main>
</body>
</html>
"""
    path.write_text(document, encoding="utf-8")


def _write_csv(result: AuditResult, path: Path) -> None:
    summary = result.to_dict()["summary"]
    rows: list[list[Any]] = [
        [
            "汇总",
            result.status,
            (
                f"实体={summary['entities']}；通过={summary['passed_entities']}；"
                f"需复核={summary['review_entities']}；错误={summary['errors']}；"
                f"警告={summary['warnings']}；输入文件={'未改动' if not result.modified_input_files else '检测到改动'}"
            ),
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    ]
    for issue in result.issue_list:
        if issue.evidence:
            for item in issue.evidence:
                rows.append(
                    [
                        "差异",
                        result.status,
                        issue.entity_id,
                        issue.severity,
                        issue.code,
                        issue.field_name or "",
                        issue.message,
                        issue.expected,
                        issue.actual,
                        item.file,
                        item.location,
                        item.excerpt,
                    ]
                )
        else:
            rows.append(
                [
                    "差异",
                    result.status,
                    issue.entity_id,
                    issue.severity,
                    issue.code,
                    issue.field_name or "",
                    issue.message,
                    issue.expected,
                    issue.actual,
                    "",
                    "",
                    "",
                ]
            )
    _safe_csv(
        path,
        [
            "记录类型",
            "总体状态",
            "实体编号",
            "问题级别",
            "问题代码",
            "字段",
            "问题说明",
            "期望",
            "实际",
            "来源文件",
            "位置",
            "原文摘录",
        ],
        rows,
    )


def _set_word_cell_text(cell: Any, text: str, bold: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text or ""))
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)


def _write_word(result: AuditResult, path: Path) -> None:
    document = WordDocument()
    section = document.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)

    normal = document.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10)

    title = document.add_heading(f"{result.profile_name}差异汇总", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        f"核验对象：{result.entity_name}　总体状态：{result.status}"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.color.rgb = RGBColor(21, 111, 115)
    subtitle.runs[0].bold = True

    summary = result.to_dict()["summary"]
    document.add_heading("一、总体汇总", level=1)
    summary_rows = [
        ("来源数量", str(summary["sources"])),
        ("文件数量", str(summary["files"])),
        ("关联实体", str(summary["entities"])),
        ("通过实体", str(summary["passed_entities"])),
        ("需复核实体", str(summary["review_entities"])),
        ("错误 / 警告", f"{summary['errors']} / {summary['warnings']}"),
        (
            "源文件完整性",
            "未改动" if not result.modified_input_files else "检测到改动",
        ),
        ("生成时间", result.finished_at),
    ]
    summary_table = document.add_table(rows=1, cols=2)
    summary_table.style = "Table Grid"
    _set_word_cell_text(summary_table.rows[0].cells[0], "项目", True)
    _set_word_cell_text(summary_table.rows[0].cells[1], "内容", True)
    for label, value in summary_rows:
        cells = summary_table.add_row().cells
        _set_word_cell_text(cells[0], label, True)
        _set_word_cell_text(cells[1], value)

    document.add_heading("二、实体对照", level=1)
    entity_table = document.add_table(rows=1, cols=4)
    entity_table.style = "Table Grid"
    for index, label in enumerate(("实体编号", "状态", "已覆盖来源", "问题数")):
        _set_word_cell_text(entity_table.rows[0].cells[index], label, True)
    for entity in result.entity_summaries:
        cells = entity_table.add_row().cells
        _set_word_cell_text(cells[0], entity["entity_id"])
        _set_word_cell_text(cells[1], entity["status"])
        _set_word_cell_text(cells[2], "、".join(entity["sources"]))
        _set_word_cell_text(cells[3], str(entity["issue_count"]))

    labels: dict[str, str] = {}
    for entity in result.entity_summaries:
        labels.update(entity.get("labels", {}))
    document.add_heading("三、差异明细", level=1)
    if result.issue_list:
        issue_table = document.add_table(rows=1, cols=6)
        issue_table.style = "Table Grid"
        for index, label in enumerate(
            ("级别", "实体", "字段", "问题", "期望", "实际")
        ):
            _set_word_cell_text(issue_table.rows[0].cells[index], label, True)
        for issue in result.issue_list:
            cells = issue_table.add_row().cells
            _set_word_cell_text(cells[0], issue.severity)
            _set_word_cell_text(cells[1], issue.entity_id)
            _set_word_cell_text(cells[2], labels.get(issue.field_name or "", issue.field_name or ""))
            _set_word_cell_text(cells[3], issue.message)
            _set_word_cell_text(cells[4], issue.expected)
            _set_word_cell_text(cells[5], issue.actual)
    else:
        document.add_paragraph("未发现配置规则范围内的差异。")

    document.add_heading("四、证据索引", level=1)
    if result.issue_list:
        for index, issue in enumerate(result.issue_list, start=1):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.add_run(f"{issue.entity_id} / {issue.code}：").bold = True
            paragraph.add_run(issue.message)
            for item in issue.evidence:
                evidence = document.add_paragraph()
                evidence.paragraph_format.left_indent = Cm(0.6)
                evidence.add_run(f"{item.file} / {item.location}：").bold = True
                evidence.add_run(item.excerpt)
    else:
        document.add_paragraph("无问题证据。")

    document.add_heading("五、来源清单", level=1)
    source_table = document.add_table(rows=1, cols=4)
    source_table.style = "Table Grid"
    for index, label in enumerate(("来源", "文件", "解析状态", "记录数")):
        _set_word_cell_text(source_table.rows[0].cells[index], label, True)
    for item in result.source_inventory:
        cells = source_table.add_row().cells
        _set_word_cell_text(cells[0], item.source)
        _set_word_cell_text(cells[1], item.file)
        _set_word_cell_text(cells[2], item.status)
        _set_word_cell_text(cells[3], str(item.records))

    document.add_paragraph()
    note = document.add_paragraph(
        "数据边界：原始资料只读；解析器仅处理临时字节副本；工具不联网、不上传资料。"
        "差异结果不等同于业务、财务或法律结论，应由责任人结合原件复核。"
    )
    note.runs[0].font.size = Pt(8)
    note.runs[0].font.color.rgb = RGBColor(96, 112, 128)

    document.core_properties.title = f"{result.profile_name}差异汇总"
    document.core_properties.subject = "多源资料一致性核验"
    document.core_properties.author = "multi-source-consistency-auditor"
    document.save(path)


def _pdf_paragraph(value: Any, style: Any) -> Paragraph:
    text = escape(str(value or "")).replace("\n", "<br/>")
    return Paragraph(text, style)


def _write_pdf(result: AuditResult, path: Path) -> None:
    bundled_font = Path(__file__).resolve().parent / "assets" / "NotoSansSC-VF.ttf"
    override_value = os.environ.get("MSCA_PDF_FONT")
    override_font = Path(override_value) if override_value else None
    if override_font and override_font.is_file():
        pdf_font_name = "MSCACustomFont"
        pdfmetrics.registerFont(TTFont(pdf_font_name, str(override_font)))
    elif bundled_font.is_file():
        pdf_font_name = "MSCASans"
        pdfmetrics.registerFont(TTFont(pdf_font_name, str(bundled_font)))
    else:
        pdf_font_name = "STSong-Light"
        pdfmetrics.registerFont(UnicodeCIDFont(pdf_font_name))

    base_styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ChineseTitle",
        parent=base_styles["Title"],
        fontName=pdf_font_name,
        fontSize=21,
        leading=29,
        textColor=colors.HexColor("#12384B"),
        alignment=TA_CENTER,
        spaceAfter=8,
    )
    subtitle_style = ParagraphStyle(
        "ChineseSubtitle",
        parent=base_styles["Normal"],
        fontName=pdf_font_name,
        fontSize=10,
        leading=16,
        textColor=colors.HexColor("#156F73"),
        alignment=TA_CENTER,
    )
    heading_style = ParagraphStyle(
        "ChineseHeading",
        parent=base_styles["Heading2"],
        fontName=pdf_font_name,
        fontSize=13,
        leading=19,
        textColor=colors.HexColor("#12384B"),
        spaceBefore=11,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        "ChineseBody",
        parent=base_styles["BodyText"],
        fontName=pdf_font_name,
        fontSize=8.5,
        leading=13,
        textColor=colors.HexColor("#17212B"),
        alignment=TA_LEFT,
        wordWrap="CJK",
    )
    small_style = ParagraphStyle(
        "ChineseSmall",
        parent=body_style,
        fontSize=7.4,
        leading=11,
        textColor=colors.HexColor("#50616F"),
        wordWrap="CJK",
    )
    cell_style = ParagraphStyle(
        "ChineseCell",
        parent=body_style,
        fontSize=7.2,
        leading=10.5,
        wordWrap="CJK",
    )
    cell_bold_style = ParagraphStyle(
        "ChineseCellBold",
        parent=cell_style,
        fontName=pdf_font_name,
        textColor=colors.white,
    )

    summary = result.to_dict()["summary"]
    labels: dict[str, str] = {}
    for entity in result.entity_summaries:
        labels.update(entity.get("labels", {}))

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#DBE3EA"))
        canvas.line(16 * mm, 13 * mm, A4[0] - 16 * mm, 13 * mm)
        canvas.setFont(pdf_font_name, 7)
        canvas.setFillColor(colors.HexColor("#607080"))
        canvas.drawString(16 * mm, 8.5 * mm, "多源资料一致性核验 · 原始资料只读且未上传")
        canvas.drawRightString(A4[0] - 16 * mm, 8.5 * mm, f"第 {doc.page} 页")
        canvas.restoreState()

    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=15 * mm,
        leftMargin=15 * mm,
        topMargin=15 * mm,
        bottomMargin=20 * mm,
        title=f"{result.profile_name}差异汇总",
        author="multi-source-consistency-auditor",
        subject="多源资料一致性核验",
    )
    story: list[Any] = [
        Paragraph(escape(f"{result.profile_name}差异汇总"), title_style),
        Paragraph(
            escape(f"核验对象：{result.entity_name}　总体状态：{result.status}"),
            subtitle_style,
        ),
        Spacer(1, 5 * mm),
        Paragraph("一、总体汇总", heading_style),
    ]

    summary_data = [
        [_pdf_paragraph("项目", cell_bold_style), _pdf_paragraph("内容", cell_bold_style)],
        ["来源数量", str(summary["sources"])],
        ["文件数量", str(summary["files"])],
        ["关联实体", str(summary["entities"])],
        ["通过实体", str(summary["passed_entities"])],
        ["需复核实体", str(summary["review_entities"])],
        ["错误 / 警告", f"{summary['errors']} / {summary['warnings']}"],
        [
            "源文件完整性",
            "未改动" if not result.modified_input_files else "检测到改动",
        ],
        ["生成时间", result.finished_at],
    ]
    summary_table = Table(summary_data, colWidths=[38 * mm, 140 * mm], repeatRows=1)
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                ("TEXTCOLOR", (0, 1), (-1, -1), colors.HexColor("#17212B")),
                ("FONTNAME", (0, 0), (-1, -1), pdf_font_name),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD6DF")),
                ("BACKGROUND", (0, 1), (0, -1), colors.HexColor("#F4F7F9")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(summary_table)
    story.append(Paragraph("二、实体对照", heading_style))

    entity_data = [
        [
            _pdf_paragraph("实体编号", cell_bold_style),
            _pdf_paragraph("状态", cell_bold_style),
            _pdf_paragraph("已覆盖来源", cell_bold_style),
            _pdf_paragraph("问题数", cell_bold_style),
        ]
    ]
    for entity in result.entity_summaries:
        entity_data.append(
            [
                _pdf_paragraph(entity["entity_id"], body_style),
                _pdf_paragraph(entity["status"], body_style),
                _pdf_paragraph("、".join(entity["sources"]), body_style),
                _pdf_paragraph(entity["issue_count"], body_style),
            ]
        )
    entity_table = Table(
        entity_data,
        colWidths=[38 * mm, 20 * mm, 94 * mm, 26 * mm],
        repeatRows=1,
    )
    entity_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD6DF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.append(entity_table)
    story.append(Paragraph("三、差异明细", heading_style))

    if result.issue_list:
        issue_data = [
            [
                _pdf_paragraph("级别", cell_bold_style),
                _pdf_paragraph("实体 / 字段", cell_bold_style),
                _pdf_paragraph("问题", cell_bold_style),
                _pdf_paragraph("期望 / 实际", cell_bold_style),
            ]
        ]
        for issue in result.issue_list:
            issue_data.append(
                [
                    _pdf_paragraph(issue.severity, cell_style),
                    _pdf_paragraph(
                        f"{issue.entity_id}\n{labels.get(issue.field_name or '', issue.field_name or '')}",
                        cell_style,
                    ),
                    _pdf_paragraph(issue.message, cell_style),
                    _pdf_paragraph(
                        f"期望：{issue.expected or '未指定'}\n实际：{issue.actual or '未指定'}",
                        cell_style,
                    ),
                ]
            )
        issue_table = Table(
            issue_data,
            colWidths=[15 * mm, 38 * mm, 55 * mm, 70 * mm],
            repeatRows=1,
        )
        issue_table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD6DF")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#FFF9EE")]),
                    ("LEFTPADDING", (0, 0), (-1, -1), 4),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(issue_table)
    else:
        story.append(Paragraph("未发现配置规则范围内的差异。", body_style))

    if result.issue_list:
        story.append(PageBreak())
        story.append(Paragraph("四、证据索引", heading_style))
        for index, issue in enumerate(result.issue_list, start=1):
            issue_heading = Paragraph(
                f"{index}. {escape(issue.entity_id)} / {escape(issue.code)}",
                body_style,
            )
            evidence_flow: list[Any] = [issue_heading]
            if issue.evidence:
                evidence_rows = []
                for item in issue.evidence:
                    evidence_rows.append(
                        [
                            _pdf_paragraph(item.file, small_style),
                            _pdf_paragraph(item.location, small_style),
                            _pdf_paragraph(item.excerpt, small_style),
                        ]
                    )
                evidence_table = Table(
                    evidence_rows,
                    colWidths=[55 * mm, 27 * mm, 96 * mm],
                    repeatRows=0,
                )
                evidence_table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9E2E8")),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                            ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F4F7F9")),
                            ("LEFTPADDING", (0, 0), (-1, -1), 4),
                            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                            ("TOPPADDING", (0, 0), (-1, -1), 3),
                            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                        ]
                    )
                )
                evidence_flow.extend([Spacer(1, 2 * mm), evidence_table, Spacer(1, 3 * mm)])
            else:
                evidence_flow.append(Paragraph("无附加证据。", small_style))
            story.append(KeepTogether(evidence_flow))
    else:
        story.append(Paragraph("四、证据索引", heading_style))
        story.append(Paragraph("无问题证据。", body_style))

    story.append(Paragraph("来源清单", heading_style))
    source_data = [
        [
            _pdf_paragraph("来源", cell_bold_style),
            _pdf_paragraph("文件", cell_bold_style),
            _pdf_paragraph("解析状态", cell_bold_style),
            _pdf_paragraph("记录数", cell_bold_style),
        ]
    ]
    for item in result.source_inventory:
        source_data.append(
            [
                _pdf_paragraph(item.source, cell_style),
                _pdf_paragraph(item.file, cell_style),
                _pdf_paragraph(item.status, cell_style),
                _pdf_paragraph(item.records, cell_style),
            ]
        )
    source_table = Table(
        source_data,
        colWidths=[31 * mm, 95 * mm, 28 * mm, 24 * mm],
        repeatRows=1,
    )
    source_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#16324F")),
                ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD6DF")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]
        )
    )
    story.extend([source_table, Spacer(1, 5 * mm)])
    story.append(
        Paragraph(
            "数据边界：原始资料只读；解析器仅处理临时字节副本；工具不联网、不上传资料。"
            "差异结果不等同于业务、财务或法律结论，应由责任人结合原件复核。",
            small_style,
        )
    )
    document.build(story, onFirstPage=footer, onLaterPages=footer)


SUPPORTED_FORMATS = {"html", "xlsx", "markdown", "json", "csv", "docx", "pdf"}


def write_reports(
    result: AuditResult,
    output_dir: str | Path,
    formats: Iterable[str] | None = None,
) -> dict[str, Path]:
    out = Path(output_dir).resolve()
    out.mkdir(parents=True, exist_ok=True)
    selected = list(formats or ["html"])
    if "all" in selected:
        selected = sorted(SUPPORTED_FORMATS)
    unknown = set(selected) - SUPPORTED_FORMATS
    if unknown:
        raise ValueError(f"不支持的输出格式：{', '.join(sorted(unknown))}")
    selected = list(dict.fromkeys(selected))

    payload = result.to_dict()
    payload["generator"] = {
        "name": "multi-source-consistency-auditor",
        "version": __version__,
    }
    paths: dict[str, Path] = {}
    if "html" in selected:
        paths["html"] = out / "differences-summary.html"
        _write_html(result, paths["html"])
    if "xlsx" in selected:
        paths["xlsx"] = out / "differences-summary.xlsx"
        _write_workbook(result, paths["xlsx"])
    if "markdown" in selected:
        paths["markdown"] = out / "differences-summary.md"
        _write_markdown(result, paths["markdown"])
    if "json" in selected:
        paths["json"] = out / "differences-summary.json"
        paths["json"].write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
    if "csv" in selected:
        paths["csv"] = out / "differences-summary.csv"
        _write_csv(result, paths["csv"])
    if "docx" in selected:
        paths["docx"] = out / "differences-summary.docx"
        _write_word(result, paths["docx"])
    if "pdf" in selected:
        paths["pdf"] = out / "differences-summary.pdf"
        _write_pdf(result, paths["pdf"])
    return paths
