"""Parse tabular and document sources into a common evidence-bearing model."""

from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import tempfile
import unicodedata
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

from .model import Evidence, Fact, Issue, Record, SourceInventory
from .normalize import (
    NormalizationError,
    normalize_identifier,
    normalize_value,
)
from .profile import FieldSpec, Profile, SourceSpec

TABLE_SUFFIXES = {".csv", ".tsv", ".xlsx", ".xlsm", ".json"}
DOCUMENT_SUFFIXES = {".docx", ".pdf", ".txt", ".md", ".markdown", ".html", ".htm"}


def _label_key(value: str) -> str:
    text = unicodedata.normalize("NFKC", str(value or ""))
    return re.sub(r"[\s:：()（）_\-]+", "", text).casefold()


def _excerpt(value: str, limit: int = 240) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text if len(text) <= limit else text[: limit - 1] + "…"


def _relative_path(path: Path, base_dir: Path) -> str:
    try:
        return str(path.resolve().relative_to(base_dir.resolve()))
    except ValueError:
        return str(path.resolve())


def _find_column(headers: list[str], aliases: Iterable[str]) -> int | None:
    normalized_headers = [_label_key(item) for item in headers]
    normalized_aliases = [_label_key(item) for item in aliases]
    for alias in normalized_aliases:
        if alias in normalized_headers:
            return normalized_headers.index(alias)
    for alias in normalized_aliases:
        for index, header in enumerate(normalized_headers):
            if alias and (alias in header or header in alias):
                return index
    return None


def _header_values(row: Iterable[Any]) -> list[str]:
    return ["" if value is None else str(value).strip() for value in row]


def _add_value(
    record: Record,
    spec: FieldSpec,
    raw_value: Any,
    evidence: Evidence,
    issues: list[Issue],
) -> None:
    if raw_value is None:
        return
    raw_text = str(raw_value).strip()
    if not raw_text:
        return
    try:
        normalized = normalize_value(spec.field_type, raw_text)
    except NormalizationError as exc:
        issues.append(
            Issue(
                severity="warning",
                code="value_normalization_failed",
                entity_id=record.record_id,
                field_name=spec.name,
                message=f"{spec.label}无法标准化：{exc}",
                actual=raw_text,
                evidence=[evidence],
            )
        )
        return
    record.add_fact(
        Fact(
            field=spec.name,
            raw_value=raw_text,
            normalized_value=normalized,
            evidence=evidence,
        )
    )


def _record_from_mapping(
    mapping: dict[str, Any],
    profile: Profile,
    source: SourceSpec,
    file_label: str,
    location: str,
    issues: list[Issue],
) -> Record:
    def column_value(value: Any) -> str:
        if value is None:
            return ""
        if hasattr(value, "isoformat"):
            try:
                return value.isoformat()
            except (TypeError, ValueError):
                pass
        return str(value).strip()

    headers = list(mapping.keys())
    id_column = _find_column(headers, profile.identifier_aliases)
    raw_id = column_value(mapping.get(headers[id_column])) if id_column is not None else ""
    record_id = normalize_identifier(raw_id)
    if not record_id:
        record_id = f"未识别ID::{source.name}::{Path(file_label).stem}::{location}"
        issues.append(
            Issue(
                severity="error",
                code="missing_identifier",
                entity_id=record_id,
                message=(
                    f"{profile.entity_name}编号为空；可尝试补充 identifier.aliases "
                    "或修正数据结构"
                ),
                expected=f"提供{profile.identifier_field}",
                actual="空值",
                evidence=[
                    Evidence(
                        source=source.name,
                        file=file_label,
                        location=location,
                        excerpt=_excerpt(" | ".join(f"{k}={v}" for k, v in mapping.items())),
                    )
                ],
            )
        )

    record = Record(
        record_id=record_id,
        source=source.name,
        file=file_label,
        location=location,
        metadata={"raw_identifier": raw_id},
    )

    for spec in profile.fields.values():
        column = _find_column(headers, spec.aliases)
        if column is None:
            continue
        raw = column_value(mapping.get(headers[column]))
        if not raw:
            continue
        evidence = Evidence(
            source=source.name,
            file=file_label,
            location=f"{location} / 列“{headers[column]}”",
            excerpt=_excerpt(f"{headers[column]}：{raw}"),
        )
        _add_value(record, spec, raw, evidence, issues)
    return record


def _read_csv_rows(path: Path) -> list[dict[str, str]]:
    errors: list[UnicodeError] = []
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                sample = handle.read(4096)
                handle.seek(0)
                delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
                if path.suffix.lower() == ".csv":
                    try:
                        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
                    except csv.Error:
                        delimiter = ","
                reader = csv.DictReader(handle, delimiter=delimiter)
                return [dict(row) for row in reader]
        except UnicodeDecodeError as exc:
            errors.append(exc)
    raise ValueError(f"CSV 编码无法识别：{errors[-1] if errors else 'unknown'}")


def _parse_table_file(
    path: Path,
    profile: Profile,
    source: SourceSpec,
    file_label: str,
    issues: list[Issue],
) -> list[Record]:
    suffix = path.suffix.lower()
    records: list[Record] = []
    if suffix in {".csv", ".tsv"}:
        for row_index, row in enumerate(_read_csv_rows(path), start=2):
            if not any(str(value or "").strip() for value in row.values()):
                continue
            records.append(
                _record_from_mapping(
                    row,
                    profile,
                    source,
                    file_label,
                    f"第 {row_index} 行",
                    issues,
                )
            )
        return records

    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook.worksheets:
                rows = list(sheet.iter_rows(values_only=True))
                header_index = next(
                    (
                        index
                        for index, row in enumerate(rows)
                        if sum(bool(str(value or "").strip()) for value in row) >= 2
                    ),
                    None,
                )
                if header_index is None:
                    continue
                headers = _header_values(rows[header_index])
                for row_offset, row in enumerate(rows[header_index + 1 :], start=header_index + 2):
                    if not any(str(value or "").strip() for value in row):
                        continue
                    padded = list(row) + [None] * max(0, len(headers) - len(row))
                    mapping = {
                        header or f"未命名列{index + 1}": padded[index]
                        for index, header in enumerate(headers)
                    }
                    records.append(
                        _record_from_mapping(
                            mapping,
                            profile,
                            source,
                            file_label,
                            f"{sheet.title} / 第 {row_offset} 行",
                            issues,
                        )
                    )
        finally:
            workbook.close()
        return records

    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8-sig"))
        rows = data if isinstance(data, list) else data.get("records", []) if isinstance(data, dict) else []
        if not isinstance(rows, list):
            raise ValueError("JSON 表格必须是对象数组，或包含 records 数组")
        for index, row in enumerate(rows, start=1):
            if not isinstance(row, dict):
                raise ValueError("JSON 表格中的每条记录都必须是对象")
            records.append(
                _record_from_mapping(
                    row,
                    profile,
                    source,
                    file_label,
                    f"记录 {index}",
                    issues,
                )
            )
        return records

    raise ValueError(f"不支持的表格格式：{suffix}")


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"p", "div", "br", "li", "tr", "h1", "h2", "h3", "h4", "table"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if data.strip():
            self.parts.append(data.strip())

    def text(self) -> str:
        return "\n".join(
            line.strip() for line in "".join(self.parts).splitlines() if line.strip()
        )


def _document_chunks(path: Path) -> list[tuple[str, str]]:
    suffix = path.suffix.lower()
    chunks: list[tuple[str, str]] = []
    if suffix == ".docx":
        document = Document(path)
        for index, paragraph in enumerate(document.paragraphs, start=1):
            if paragraph.text.strip():
                chunks.append((f"段落 {index}", paragraph.text.strip()))
        for table_index, table in enumerate(document.tables, start=1):
            for row_index, row in enumerate(table.rows, start=1):
                text = " | ".join(cell.text.strip() for cell in row.cells)
                if text.strip(" |"):
                    chunks.append((f"表格 {table_index} / 第 {row_index} 行", text))
        return chunks

    if suffix == ".pdf":
        reader = PdfReader(str(path))
        for page_index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                chunks.append((f"第 {page_index} 页", text.strip()))
        return chunks

    if suffix in {".txt", ".md", ".markdown"}:
        try:
            text = path.read_text(encoding="utf-8-sig")
        except UnicodeDecodeError:
            text = path.read_text(encoding="gb18030")
        for index, line in enumerate(text.splitlines(), start=1):
            if line.strip():
                chunks.append((f"第 {index} 行", line.strip()))
        return chunks

    if suffix in {".html", ".htm"}:
        extractor = _HTMLTextExtractor()
        extractor.feed(path.read_text(encoding="utf-8-sig"))
        for index, line in enumerate(extractor.text().splitlines(), start=1):
            chunks.append((f"文本块 {index}", line))
        return chunks

    raise ValueError(f"不支持的文档格式：{suffix}")


def _find_id(
    profile: Profile,
    path: Path,
    chunks: list[tuple[str, str]],
) -> tuple[str, str, str]:
    if profile.identifier_pattern:
        pattern = re.compile(profile.identifier_pattern, re.IGNORECASE)
        filename_match = pattern.search(path.stem)
        if filename_match:
            raw = filename_match.group(0)
            return normalize_identifier(raw), f"文件名“{path.name}”", raw
        for location, text in chunks:
            content_match = pattern.search(text)
            if content_match:
                raw = content_match.group(0)
                return normalize_identifier(raw), location, raw

    for location, text in chunks:
        for segment in re.split(r"[\n\r；;|]+", text):
            for alias in profile.identifier_aliases:
                pattern = re.compile(
                    rf"(?:^|[\s，,])(?:{re.escape(alias)})\s*[:：]\s*([^\n\r；;|]+)",
                    re.IGNORECASE,
                )
                match = pattern.search(segment)
                if match:
                    raw = match.group(1).strip()
                    return normalize_identifier(raw), location, raw
    return "", "", ""


def _find_alias_value(text: str, aliases: Iterable[str]) -> str:
    for alias in aliases:
        pattern = re.compile(
            rf"(?:^|[\n\r；;|，,])\s*{re.escape(alias)}\s*[:：]\s*([^\n\r；;|]+)",
            re.IGNORECASE,
        )
        match = pattern.search(text)
        if match:
            return match.group(1).strip()
    return ""


def _parse_document_file(
    path: Path,
    profile: Profile,
    source: SourceSpec,
    file_label: str,
    issues: list[Issue],
) -> list[Record]:
    chunks = _document_chunks(path)
    if not chunks:
        issues.append(
            Issue(
                severity="error",
                code="empty_or_scanned_document",
                entity_id=f"文件::{Path(file_label).stem}",
                message="未提取到文本；文件可能为空、损坏或为扫描件",
                evidence=[
                    Evidence(
                        source=source.name,
                        file=file_label,
                        location="整个文件",
                        excerpt="无可用文本",
                    )
                ],
            )
        )
        return []

    record_id, id_location, raw_id = _find_id(profile, path, chunks)
    if not record_id:
        record_id = f"未识别ID::{source.name}::{Path(file_label).stem}"
        issues.append(
            Issue(
                severity="error",
                code="missing_identifier",
                entity_id=record_id,
                message="未能从文件名或正文识别实体编号",
                expected=f"提供{profile.identifier_field}或设置 identifier.pattern",
                actual="空值",
                evidence=[
                    Evidence(
                        source=source.name,
                        file=file_label,
                        location="文件名及正文",
                        excerpt=_excerpt(Path(file_label).name),
                    )
                ],
            )
        )

    record = Record(
        record_id=record_id,
        source=source.name,
        file=file_label,
        location=id_location or "文件",
        metadata={"raw_identifier": raw_id},
    )
    for spec in profile.fields.values():
        for location, text in chunks:
            raw = _find_alias_value(text, spec.aliases)
            if not raw:
                continue
            evidence = Evidence(
                source=source.name,
                file=file_label,
                location=location,
                excerpt=_excerpt(text),
            )
            _add_value(record, spec, raw, evidence, issues)
            break
    return [record]


def _expand_source_files(profile: Profile, source: SourceSpec) -> tuple[list[Path], list[Issue]]:
    files: list[Path] = []
    issues: list[Issue] = []
    for path in profile.resolve_paths(source):
        if path.is_dir():
            candidates = [
                item
                for item in path.rglob("*")
                if item.is_file() and not item.name.startswith("~$")
            ]
            if source.include:
                candidates = [
                    item
                    for item in candidates
                    if any(item.match(pattern) for pattern in source.include)
                ]
            files.extend(candidates)
        elif path.is_file():
            files.append(path)
    unique = sorted({item.resolve() for item in files})
    if not unique:
        issues.append(
            Issue(
                severity="error",
                code="source_not_found",
                entity_id="来源",
                message=f"来源“{source.name}”未找到任何文件",
                actual=", ".join(source.paths),
            )
        )
    return unique, issues


def snapshot_source_files(profile: Profile) -> dict[Path, str]:
    """Return SHA-256 hashes without reading or modifying source content elsewhere."""
    snapshots: dict[Path, str] = {}
    for source in profile.sources:
        files, _issues = _expand_source_files(profile, source)
        for path in files:
            digest = hashlib.sha256()
            with path.open("rb") as handle:
                for block in iter(lambda: handle.read(1024 * 1024), b""):
                    digest.update(block)
            snapshots[path.resolve()] = digest.hexdigest()
    return snapshots


def collect_records(
    profile: Profile,
) -> tuple[list[Record], list[SourceInventory], list[Issue]]:
    records: list[Record] = []
    inventory: list[SourceInventory] = []
    issues: list[Issue] = []

    for source in profile.sources:
        files, source_issues = _expand_source_files(profile, source)
        issues.extend(source_issues)
        for path in files:
            file_label = _relative_path(path, profile.base_dir)
            suffix = path.suffix.lower()
            source_type = (
                "table"
                if suffix in TABLE_SUFFIXES
                else "document"
                if suffix in DOCUMENT_SUFFIXES
                else "unsupported"
            )
            if source.kind != "auto" and source.kind != source_type:
                issues.append(
                    Issue(
                        severity="error",
                        code="source_type_mismatch",
                        entity_id="来源",
                        message=(
                            f"来源“{source.name}”配置为 {source.kind}，"
                            f"但文件“{file_label}”识别为 {source_type}"
                        ),
                    )
                )
                inventory.append(
                    SourceInventory(source.name, file_label, source_type, "已跳过", 0)
                )
                continue

            try:
                if source_type == "table":
                    parser = _parse_table_file
                elif source_type == "document":
                    parser = _parse_document_file
                else:
                    raise ValueError(f"不支持的文件格式：{suffix or '无扩展名'}")
                # Third-party parsers only receive a disposable copy. The original file
                # is opened read-only for hashing and byte copying, never for parsing
                # or writing.
                with tempfile.TemporaryDirectory(prefix="msca-readonly-") as temp_dir:
                    temporary_copy = Path(temp_dir) / path.name
                    shutil.copyfile(path, temporary_copy)
                    parsed = parser(
                        temporary_copy,
                        profile,
                        source,
                        file_label,
                        issues,
                    )
                records.extend(parsed)
                inventory.append(
                    SourceInventory(source.name, file_label, source_type, "已解析", len(parsed))
                )
            except Exception as exc:  # A single bad file must not abort the whole audit.
                issues.append(
                    Issue(
                        severity="error",
                        code="file_parse_failed",
                        entity_id="来源",
                        message=f"文件解析失败：{exc}",
                        evidence=[
                            Evidence(
                                source=source.name,
                                file=file_label,
                                location="整个文件",
                                excerpt=_excerpt(str(exc)),
                            )
                        ],
                    )
                )
                inventory.append(
                    SourceInventory(source.name, file_label, source_type, "解析失败", 0)
                )

    return records, inventory, issues
